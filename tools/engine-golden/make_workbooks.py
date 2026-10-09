"""Writes .xlsx files with every kind of cell the identifier check meets, for the parity tests
(ParityTests) to scan with both engines. Synthetic values only.

    uv run --project <LabOps-Projects clone> python tools/engine-golden/make_workbooks.py <folder>
"""
import datetime as dt
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont


def openpyxl_book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Samples"
    ws.append(["Sample_ID", "Collected", "Time", "Duration", "Weight", "Count", "OK", "Formula", "Patient Name", "Notes"])
    rows = [
        ["S1", dt.datetime(2026, 1, 2), dt.time(9, 30), dt.timedelta(hours=26, minutes=3), 61.5, 3, True, "=E2*2", "Ann", "a b c d e f"],
        ["S2", dt.datetime(2026, 1, 2, 14, 5, 7), dt.time(0, 0, 1), dt.timedelta(seconds=1.5), 70.0, 0, False, "=SUM(F2:F3)", None, None],
        ["S3", dt.date(2025, 12, 31), None, None, 1e20, -12, None, None, "Bob", "call 206-555-0100"],
    ]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "yyyy-mm-dd h:mm"
    ws["B4"].number_format = "d-mmm-yy"
    ws["C2"].number_format = "h:mm"
    ws["C3"].number_format = "h:mm:ss"
    ws["D2"].number_format = "[h]:mm:ss"
    ws["D3"].number_format = "[mm]:ss.0"
    ws["E2"].number_format = '0.00" kg"'
    ws["E3"].number_format = '[Red]0.0;[Blue]-0.0'
    ws["F2"].number_format = '"Day "0'
    ws["K1"] = "Rich"
    ws["K2"] = CellRichText([TextBlock(InlineFont(b=True), "bold"), " and plain"])
    ws["Z10"].number_format = "0"    # a styled empty cell far from the data
    second = wb.create_sheet("Hidden")
    second.sheet_state = "hidden"
    second.append(["Email", "Value"])
    second.append(["x@y.org", 1])
    wb.create_sheet("Empty")
    wb.save(path)


CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
<Override PartName="/xl/chartsheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.chartsheet+xml"/>
<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>
<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
</Types>"""

ROOT_RELS = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>"""

WORKBOOK = """<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<workbookPr date1904="1"/>
<sheets>
<sheet name="Data" sheetId="1" r:id="rId1"/>
<sheet name="Chart" sheetId="2" r:id="rId3"/>
<sheet name="More" sheetId="3" state="hidden" r:id="rId2"/>
</sheets>
</workbook>"""

WORKBOOK_RELS = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="/xl/worksheets/sheet2.xml"/>
<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chartsheet" Target="chartsheets/sheet1.xml"/>
<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>
<Relationship Id="rId5" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>"""

STRINGS = """<?xml version="1.0" encoding="UTF-8"?>
<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="6" uniqueCount="6">
<si><t>Subject</t></si>
<si><r><rPr><b/></rPr><t>Home</t></r><r><t xml:space="preserve"> address</t></r></si>
<si><t xml:space="preserve">  padded  </t></si>
<si><t>DOB</t><rPh sb="0" eb="1"><t>ignored</t></rPh></si>
<si><t>a_x005F_x000D_b</t></si>
<si><t/></si>
</sst>"""

STYLES = """<?xml version="1.0" encoding="UTF-8"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<numFmts count="2"><numFmt numFmtId="164" formatCode="yyyy\\-mm\\-dd"/><numFmt numFmtId="165" formatCode="&quot;Lot &quot;0"/></numFmts>
<cellXfs count="6">
<xf numFmtId="0"/><xf numFmtId="14"/><xf numFmtId="164"/><xf numFmtId="165"/><xf numFmtId="21"/><xf numFmtId="46"/>
</cellXfs>
</styleSheet>"""

SHEET1 = """<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<dimension ref="A1:B2"/>
<sheetData>
<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c><c r="C1" t="s"><v>3</v></c><c r="D1" t="inlineStr"><is><t>Inline</t></is></c><c r="E1" t="s"><v>4</v></c><c r="F1" t="s"><v>5</v></c></row>
<row r="2"><c r="A2"><v>1</v></c><c r="B2" t="s"><v>2</v></c><c r="C2" s="1"><v>45000</v></c><c r="D2" t="inlineStr"><is><r><t>rich</t></r><r><t> inline</t></r></is></c><c r="E2" s="3"><v>7</v></c></row>
<row r="4"><c r="A4"><v>2.0</v></c><c r="B4" t="b"><v>1</v></c><c r="C4" s="2"><v>45000.75</v></c><c r="D4" t="e"><v>#N/A</v></c><c r="E4" t="str"><f>A1</f><v>Subject</v></c><c r="H4" s="4"/></row>
<row r="3"><c r="A3" t="inlineStr"><is><t>out of order</t></is></c></row>
<row><c><v>3</v></c><c t="d"><v>2026-01-02T03:04:05</v></c><c s="4"><v>0.5</v></c><c s="5"><v>1.25</v></c><c s="1"><v>59</v></c><c s="1"><v>61</v></c></row>
<row r="9" spans="1:3"><c r="B9"><v>1E3</v></c><c r="A9"><v>-4</v></c></row>
<row r="10"/>
</sheetData>
</worksheet>"""

SHEET2 = """<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>
<row r="1"><c r="A1" t="inlineStr"><is><t>Phone</t></is></c></row>
<row r="2"><c r="A2" t="inlineStr"><is><t>(206) 555-0100</t></is></c></row>
</sheetData></worksheet>"""

CHARTSHEET = """<?xml version="1.0" encoding="UTF-8"?>
<chartsheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>"""


def crafted_book(path: Path):
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", CONTENT_TYPES)
        z.writestr("_rels/.rels", ROOT_RELS)
        z.writestr("xl/workbook.xml", WORKBOOK)
        z.writestr("xl/_rels/workbook.xml.rels", WORKBOOK_RELS)
        z.writestr("xl/sharedStrings.xml", STRINGS)
        z.writestr("xl/styles.xml", STYLES)
        z.writestr("xl/worksheets/sheet1.xml", SHEET1)
        z.writestr("xl/worksheets/sheet2.xml", SHEET2)
        z.writestr("xl/chartsheets/sheet1.xml", CHARTSHEET)
        z.writestr("xl/chartsheets/_rels/sheet1.xml.rels",
                   '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>')


def main() -> int:
    folder = Path(sys.argv[1])
    folder.mkdir(parents=True, exist_ok=True)
    openpyxl_book(folder / "built.xlsx")
    crafted_book(folder / "crafted.xlsx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
